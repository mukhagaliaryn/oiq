import base64
import json
import re
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import anthropic
from django.conf import settings
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.utils.translation import gettext_lazy as _

from core.utils.files import question_import_image_path

IMAGE_PLACEHOLDER_RE = re.compile(r'\{\{\s*img\s*:\s*(\d+)\s*\}\}')

QUESTION_SCHEMA = {
    'type': 'object',
    'properties': {
        'questions': {
            'type': 'array',
            'items': {
                'type': 'object',
                'properties': {
                    'format_code': {
                        'type': 'string',
                        'enum': ['test', 'unsupported'],
                    },
                    'variant_code': {
                        'type': 'string',
                        'enum': ['single', 'multiple'],
                    },
                    'level': {
                        'type': 'string',
                        'enum': ['easy', 'medium', 'hard'],
                    },
                    'text_html': {'type': 'string'},
                    'options': {
                        'type': 'array',
                        'items': {
                            'type': 'object',
                            'properties': {
                                'text_html': {'type': 'string'},
                                'is_correct': {'type': 'boolean'},
                            },
                            'required': ['text_html', 'is_correct'],
                            'additionalProperties': False,
                        },
                    },
                    'warning': {'type': 'string'},
                },
                'required': ['format_code', 'variant_code', 'level', 'text_html', 'options', 'warning'],
                'additionalProperties': False,
            },
        },
    },
    'required': ['questions'],
    'additionalProperties': False,
}

IMPORT_PROMPT = """\
You are extracting exam questions from the attached document (rendered as a PDF) so they can be \
imported into a question bank. Read the whole document from top to bottom and return every question \
you find as a JSON object matching the provided schema. Follow these rules exactly:

1. Preserve the original wording verbatim, in whatever language it is written (Kazakh, Russian, English, ...). \
   Do not translate, summarize, or paraphrase anything.
2. Each question becomes one item in "questions", in the order they appear in the document.
3. "format_code": use "test" only for single-choice or multiple-choice questions that have a list of \
   answer options where one or more are marked/known to be correct. Use "unsupported" for anything else \
   (matching pairs, fill-in-the-blank, open-ended/essay questions, etc.) — for those, still fill in \
   "text_html" with the question text and leave "options" as an empty array, and explain briefly in \
   "warning" why it was not imported (e.g. "matching-type question, add manually").
4. "variant_code": "single" if exactly one option is correct, "multiple" if more than one option is correct. \
   For "unsupported" questions, set this to "single" as a placeholder.
5. "level": infer the difficulty (easy/medium/hard) from the question's complexity. Default to "medium" \
   if you cannot tell.
6. "text_html" and each option's "text_html" must be plain HTML fragments (no <html>/<body> wrapper). Use \
   <p>, <ul>/<ol>/<li>, and <table>/<tr>/<td> tags to preserve paragraphs, lists, and tables exactly as \
   they are laid out in the source.
7. Mathematical formulas: transcribe every formula into LaTeX and wrap it exactly like this — inline: \
   <span class="math-tex">\\(x^2+1\\)</span>, display/block: <span class="math-tex">\\[x^2+1\\]</span>. \
   Transcribe formulas with complete accuracy — do not alter exponents, coefficients, or symbols.
8. Images: number every embedded image sequentially across the whole document, starting at 1, in the \
   order it is first encountered while reading top to bottom. At the exact place an image appears \
   (inside a question's text or inside an option's text), insert a placeholder token in that exact form: \
   {{img:N}} — where N is that image's number. Do not describe the image, only place the token.
9. "warning" must be an empty string "" when there is nothing to flag.
10. Do not invent questions or options that are not in the document. Do not skip any question.
"""


@dataclass
class ParsedOption:
    text_html: str
    is_correct: bool


@dataclass
class ParsedQuestion:
    format_code: str
    variant_code: str
    level: str
    text_html: str
    options: list = field(default_factory=list)
    warning: str = ''

    @property
    def is_supported(self):
        return self.format_code == 'test'


class QuestionImportError(Exception):
    pass


def _convert_docx_to_pdf(docx_bytes):
    with tempfile.TemporaryDirectory() as tmpdir:
        docx_path = Path(tmpdir) / 'source.docx'
        docx_path.write_bytes(docx_bytes)

        profile_dir = Path(tmpdir) / 'lo_profile'
        try:
            subprocess.run(
                [
                    'soffice', '--headless', '--norestore',
                    f'-env:UserInstallation=file://{profile_dir}',
                    '--convert-to', 'pdf', '--outdir', tmpdir, str(docx_path),
                ],
                check=True, capture_output=True, timeout=120,
            )
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired, FileNotFoundError) as error:
            raise QuestionImportError(_('Could not convert the document. Please check the file and try again.')) from error

        pdf_path = docx_path.with_suffix('.pdf')
        if not pdf_path.exists():
            raise QuestionImportError(_('Could not convert the document. Please check the file and try again.'))

        return pdf_path.read_bytes()


_DRAWING_BLIP_TAG = '{http://schemas.openxmlformats.org/drawingml/2006/main}blip'
_RELATIONSHIP_EMBED_ATTR = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'


def _extract_ordered_images(docx_bytes):
    from docx import Document

    with tempfile.NamedTemporaryFile(suffix='.docx') as tmp:
        tmp.write(docx_bytes)
        tmp.flush()
        document = Document(tmp.name)

        images = []
        for blip in document.element.body.iter(_DRAWING_BLIP_TAG):
            r_id = blip.get(_RELATIONSHIP_EMBED_ATTR)
            if not r_id:
                continue

            try:
                part = document.part.related_parts[r_id]
            except KeyError:
                continue

            images.append({'blob': part.blob, 'content_type': part.content_type})

        return images


def _store_images(images, batch_id):
    urls = []
    paths = []
    for index, image in enumerate(images, start=1):
        extension = '.' + image['content_type'].split('/')[-1].split('+')[0]
        if extension == '.jpeg':
            extension = '.jpg'

        path = question_import_image_path(batch_id, index, extension)
        saved_path = default_storage.save(path, ContentFile(image['blob']))
        urls.append(default_storage.url(saved_path))
        paths.append(saved_path)

    return urls, paths


def _substitute_image_placeholders(html, image_urls):
    def replace(match):
        index = int(match.group(1))
        if 1 <= index <= len(image_urls):
            return f'<img src="{image_urls[index - 1]}" alt="">'

        return '<span class="text-danger">[' + str(_('image not found')) + ']</span>'

    return IMAGE_PLACEHOLDER_RE.sub(replace, html)


def _call_claude(pdf_bytes):
    client = anthropic.Anthropic(api_key=settings.ANTHROPIC_API_KEY)
    encoded = base64.standard_b64encode(pdf_bytes).decode('utf-8')

    try:
        with client.messages.stream(
            model=settings.QUESTION_IMPORT_MODEL,
            max_tokens=64000,
            thinking={'type': 'adaptive'},
            output_config={
                'effort': 'high',
                'format': {'type': 'json_schema', 'schema': QUESTION_SCHEMA},
            },
            messages=[{
                'role': 'user',
                'content': [
                    {
                        'type': 'document',
                        'source': {'type': 'base64', 'media_type': 'application/pdf', 'data': encoded},
                    },
                    {'type': 'text', 'text': IMPORT_PROMPT},
                ],
            }],
        ) as stream:
            response = stream.get_final_message()
    except anthropic.APIError as error:
        raise QuestionImportError(_('The AI service could not process the file. Please try again later.')) from error

    if response.stop_reason == 'refusal':
        raise QuestionImportError(_('The AI service declined to process this file.'))

    text_block = next((block.text for block in response.content if block.type == 'text'), None)
    if not text_block:
        raise QuestionImportError(_('The AI service returned an empty response.'))

    try:
        return json.loads(text_block)
    except json.JSONDecodeError as error:
        raise QuestionImportError(_('The AI service returned an unreadable response.')) from error


@dataclass
class ImportResult:
    questions: list
    image_paths: list


def run_question_import(docx_bytes, batch_id):
    pdf_bytes = _convert_docx_to_pdf(docx_bytes)
    images = _extract_ordered_images(docx_bytes)
    payload = _call_claude(pdf_bytes)

    image_urls, image_paths = _store_images(images, batch_id)

    questions = []
    for raw in payload.get('questions', []):
        text_html = _substitute_image_placeholders(raw.get('text_html', ''), image_urls)
        options = [
            ParsedOption(
                text_html=_substitute_image_placeholders(option.get('text_html', ''), image_urls),
                is_correct=bool(option.get('is_correct')),
            )
            for option in raw.get('options', [])
        ]

        questions.append(ParsedQuestion(
            format_code=raw.get('format_code', 'unsupported'),
            variant_code=raw.get('variant_code', 'single'),
            level=raw.get('level', 'medium'),
            text_html=text_html,
            options=options,
            warning=raw.get('warning', ''),
        ))

    return ImportResult(questions=questions, image_paths=image_paths)
